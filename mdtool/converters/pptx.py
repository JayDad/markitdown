from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


class PptxConverter:
    name = "pptx"
    extensions = (".pptx",)

    def convert(self, path: Path) -> str:
        prs = Presentation(str(path))
        out: list[str] = [f"# {path.stem}", ""]

        for idx, slide in enumerate(prs.slides, start=1):
            title_shape = self._title_shape(slide)
            title_text = self._first_line(title_shape) if title_shape else ""

            header = f"## Slide {idx}"
            if title_text:
                header += f": {title_text}"
            out.append(header)
            out.append("")

            for shape in slide.shapes:
                if title_shape is not None and shape is title_shape:
                    continue
                rendered = self._render_shape(shape, title_shape)
                if rendered:
                    out.append(rendered)
                    out.append("")

            notes = self._extract_notes(slide)
            if notes:
                for line in notes.splitlines():
                    out.append(f"> {line}" if line else ">")
                out.append("")

        return "\n".join(out).rstrip() + "\n"

    @staticmethod
    def _title_shape(slide):
        try:
            return slide.shapes.title
        except Exception:
            return None

    @staticmethod
    def _first_line(shape) -> str:
        if shape is None or not getattr(shape, "has_text_frame", False):
            return ""
        text = shape.text_frame.text.strip()
        if not text:
            return ""
        return text.splitlines()[0].strip()

    def _render_shape(self, shape, title_shape) -> str:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            parts = []
            for sub in shape.shapes:
                if title_shape is not None and sub is title_shape:
                    continue
                rendered = self._render_shape(sub, title_shape)
                if rendered:
                    parts.append(rendered)
            return "\n\n".join(parts)

        if getattr(shape, "has_table", False) and shape.has_table:
            return self._render_table(shape.table)

        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            return self._render_text_frame(shape.text_frame)

        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return "*[이미지]*"

        return ""

    @staticmethod
    def _render_text_frame(tf) -> str:
        lines: list[str] = []
        for p in tf.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            level = p.level or 0
            indent = "  " * level
            lines.append(f"{indent}- {text}")
        return "\n".join(lines)

    @staticmethod
    def _render_table(table) -> str:
        rows = list(table.rows)
        if not rows:
            return ""
        body: list[list[str]] = []
        for row in rows:
            cells = [cell.text.strip().replace("\n", " ").replace("|", "\\|") for cell in row.cells]
            body.append(cells)

        if not body or not body[0]:
            return ""

        header = "| " + " | ".join(body[0]) + " |"
        sep = "| " + " | ".join("---" for _ in body[0]) + " |"
        rest = ["| " + " | ".join(r) + " |" for r in body[1:]]
        return "\n".join([header, sep, *rest])

    @staticmethod
    def _extract_notes(slide) -> str:
        if not slide.has_notes_slide:
            return ""
        ntf = slide.notes_slide.notes_text_frame
        if ntf is None:
            return ""
        return ntf.text.strip()
